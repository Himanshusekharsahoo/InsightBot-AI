"""Document Parsers for InsightBot AI Knowledge Ingestion.

Supports PDF, DOCX, TXT, and Markdown documents with robust error handling,
page tracking, and metadata extraction.
"""

import io
import logging
import re
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

import docx
import pypdf
from pydantic import BaseModel, Field

from insightbot.security.policy import SensitiveDataSanitizer

logger = logging.getLogger('insightbot.knowledge.parsers')


class DocumentPage(BaseModel):
	"""Single page or section extracted from a document."""

	page_number: int = Field(..., ge=1, description='1-indexed page or section number')
	text: str = Field(..., description='Extracted text from this page')


class ParsedDocument(BaseModel):
	"""Standardized outcome of parsing a source document."""

	filename: str = Field(..., description='Original filename')
	source_type: str = Field(..., description='File format identifier: pdf, docx, txt, markdown')
	mime_type: str = Field(..., description='MIME type')
	text: str = Field(..., description='Complete extracted text')
	pages: list[DocumentPage] = Field(default_factory=list, description='Page-by-page text segments')
	metadata: dict[str, Any] = Field(default_factory=dict, description='Extracted document metadata')


class UnsupportedDocumentError(ValueError):
	"""Raised when a document type is unsupported."""

	pass


class DocumentParsingError(ValueError):
	"""Raised when parsing fails due to corruption or malformed content."""

	pass


def _sanitize_parser_error(action: str, filename: str, exc: Exception) -> str:
	"""Format actionable parser error without leaking internal filesystem paths, credentials, or raw stack traces."""
	err_str = str(exc)
	err_str = re.sub(r'[A-Za-z]:\\[^:\s,;]+', '[path]', err_str)
	err_str = re.sub(r'/(?:[a-zA-Z0-9_\-\.]+/)+[a-zA-Z0-9_\-\.]+', '[path]', err_str)
	err_str = SensitiveDataSanitizer.sanitize(err_str)
	if err_str and len(err_str) <= 120 and '\n' not in err_str:
		return f"{action} '{filename}': {err_str}"
	return f"{action} '{filename}': {type(exc).__name__}"


class BaseDocumentParser(ABC):
	"""Abstract base parser for document ingestion."""

	@abstractmethod
	def parse(self, filename: str, content: bytes) -> ParsedDocument:
		"""Extract clean text and metadata from document bytes."""
		pass


class PDFParser(BaseDocumentParser):
	"""PDF document parser using pypdf."""

	def parse(self, filename: str, content: bytes) -> ParsedDocument:
		if not content:
			return ParsedDocument(
				filename=filename,
				source_type='pdf',
				mime_type='application/pdf',
				text='',
				pages=[],
				metadata={'doc_type': 'pdf', 'page_count': 0, 'empty': True},
			)

		try:
			stream = io.BytesIO(content)
			reader = pypdf.PdfReader(stream)
		except Exception as e:
			raise DocumentParsingError(_sanitize_parser_error("Failed to parse PDF document", filename, e)) from e

		if reader.is_encrypted:
			try:
				reader.decrypt('')
			except Exception as e:
				raise DocumentParsingError(_sanitize_parser_error("PDF document is encrypted and cannot be decrypted", filename, e)) from e

		pages: list[DocumentPage] = []
		text_parts: list[str] = []

		for idx, page in enumerate(reader.pages, start=1):
			try:
				extracted = page.extract_text() or ''
			except Exception as e:
				logger.warning(f"Error extracting text from page {idx} of '{filename}': {e}")
				extracted = ''

			cleaned = extracted.strip()
			if cleaned:
				pages.append(DocumentPage(page_number=idx, text=cleaned))
				text_parts.append(cleaned)

		full_text = '\n\n'.join(text_parts)
		doc_meta: dict[str, Any] = {
			'doc_type': 'pdf',
			'page_count': len(reader.pages),
		}

		if reader.metadata:
			try:
				if reader.metadata.author:
					doc_meta['author'] = str(reader.metadata.author)
				if reader.metadata.title:
					doc_meta['title'] = str(reader.metadata.title)
			except Exception:
				pass

		return ParsedDocument(
			filename=filename,
			source_type='pdf',
			mime_type='application/pdf',
			text=full_text,
			pages=pages,
			metadata=doc_meta,
		)


class DocxParser(BaseDocumentParser):
	"""DOCX document parser using python-docx."""

	def parse(self, filename: str, content: bytes) -> ParsedDocument:
		if not content:
			return ParsedDocument(
				filename=filename,
				source_type='docx',
				mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
				text='',
				pages=[],
				metadata={'doc_type': 'docx', 'empty': True},
			)

		try:
			stream = io.BytesIO(content)
			doc = docx.Document(stream)
		except Exception as e:
			raise DocumentParsingError(_sanitize_parser_error("Failed to parse DOCX document", filename, e)) from e

		paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

		# Also extract table text
		for table in doc.tables:
			for row in table.rows:
				row_text = ' | '.join(c.text.strip() for c in row.cells if c.text.strip())
				if row_text:
					paragraphs.append(row_text)

		full_text = '\n\n'.join(paragraphs)
		page = DocumentPage(page_number=1, text=full_text) if full_text else None

		return ParsedDocument(
			filename=filename,
			source_type='docx',
			mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
			text=full_text,
			pages=[page] if page else [],
			metadata={'doc_type': 'docx', 'paragraph_count': len(paragraphs)},
		)


class TextParser(BaseDocumentParser):
	"""Plain text document parser with multi-encoding fallback."""

	def parse(self, filename: str, content: bytes) -> ParsedDocument:
		if not content:
			return ParsedDocument(
				filename=filename,
				source_type='txt',
				mime_type='text/plain',
				text='',
				pages=[],
				metadata={'doc_type': 'txt', 'empty': True},
			)

		decoded = None
		for encoding in ('utf-8', 'utf-8-sig', 'latin-1', 'cp1252'):
			try:
				decoded = content.decode(encoding)
				break
			except UnicodeDecodeError:
				continue

		if decoded is None:
			raise DocumentParsingError(f"Failed to decode text file '{filename}' with supported encodings.")

		text = decoded.strip()
		page = DocumentPage(page_number=1, text=text) if text else None

		return ParsedDocument(
			filename=filename,
			source_type='txt',
			mime_type='text/plain',
			text=text,
			pages=[page] if page else [],
			metadata={'doc_type': 'txt'},
		)


class MarkdownParser(BaseDocumentParser):
	"""Markdown document parser."""

	def parse(self, filename: str, content: bytes) -> ParsedDocument:
		if not content:
			return ParsedDocument(
				filename=filename,
				source_type='markdown',
				mime_type='text/markdown',
				text='',
				pages=[],
				metadata={'doc_type': 'markdown', 'empty': True},
			)

		decoded = None
		for encoding in ('utf-8', 'utf-8-sig', 'latin-1'):
			try:
				decoded = content.decode(encoding)
				break
			except UnicodeDecodeError:
				continue

		if decoded is None:
			raise DocumentParsingError(f"Failed to decode markdown file '{filename}'.")

		text = decoded.strip()
		page = DocumentPage(page_number=1, text=text) if text else None

		return ParsedDocument(
			filename=filename,
			source_type='markdown',
			mime_type='text/markdown',
			text=text,
			pages=[page] if page else [],
			metadata={'doc_type': 'markdown'},
		)


class DocumentParserFactory:
	"""Factory to select the appropriate parser based on file extension."""

	_PARSERS: dict[str, BaseDocumentParser] = {
		'.pdf': PDFParser(),
		'.docx': DocxParser(),
		'.txt': TextParser(),
		'.md': MarkdownParser(),
		'.markdown': MarkdownParser(),
	}

	@classmethod
	def get_parser(cls, filename: str) -> BaseDocumentParser:
		ext = Path(filename).suffix.lower()
		if ext not in cls._PARSERS:
			raise UnsupportedDocumentError(
				f"Unsupported file format '{ext}'. Supported formats: {', '.join(sorted(cls._PARSERS.keys()))}"
			)
		return cls._PARSERS[ext]

	@classmethod
	def parse_file(cls, filename: str, content: bytes) -> ParsedDocument:
		"""Parse file contents using the matching parser."""
		parser = cls.get_parser(filename)
		return parser.parse(filename=filename, content=content)


__all__ = [
	'DocumentPage',
	'ParsedDocument',
	'UnsupportedDocumentError',
	'DocumentParsingError',
	'BaseDocumentParser',
	'PDFParser',
	'DocxParser',
	'TextParser',
	'MarkdownParser',
	'DocumentParserFactory',
]
